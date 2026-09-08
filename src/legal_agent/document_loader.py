"""Load supported reference and case document formats into plain text."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any


def load_document(uploaded_file: Any) -> str:
    """Extract text from a Streamlit upload with a supported file extension."""
    suffix = Path(uploaded_file.name).suffix.lower()
    payload = uploaded_file.getvalue()

    if suffix == ".txt":
        return payload.decode("utf-8-sig")

    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(BytesIO(payload)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        text = "\n".join(pages).strip()
        if not text:
            raise ValueError(
                f"Could not extract text from {uploaded_file.name}. "
                "Scanned PDFs need OCR before upload."
            )
        return text

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(payload))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join(item for item in paragraphs + table_rows if item).strip()
        if not text:
            raise ValueError(f"Could not extract text from {uploaded_file.name}.")
        return text

    raise ValueError(
        f"Unsupported file type: {suffix or 'unknown'}. "
        "Upload a TXT, PDF, or DOCX document."
    )
